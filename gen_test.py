from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.title = "Исполняемые"
for i in range(1, 8):
    ws.cell(row=i, column=1, value=str(i))
headers = ['H'] * 34
for c,h in enumerate(headers,1):
    ws.cell(row=8, column=c, value=h)
row = [''] * 34
row[0] = '1'
row[1] = 'АО Тест'
row[2] = 'Д-001'
row[3] = 'Работа 1'
row[12] = '1000000'
row[30] = 'Приостановлено оформление'
for c,v in enumerate(row,1):
    ws.cell(row=9, column=c, value=v)
wb.save('/tmp/t.xlsx')
import pandas as pd
df = pd.read_excel('/tmp/t.xlsx', sheet_name='Исполняемые', header=None, dtype=str)
for i in range(28, 33):
    print(f'idx {i}: {repr(df.iloc[8, i])}')
